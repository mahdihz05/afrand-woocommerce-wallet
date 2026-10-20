"""Normalize a sales-invoice export without changing its financial meaning."""

import re
from datetime import date, datetime

from openpyxl import load_workbook

FIELDS = [
    ("order_ref", "شماره فاکتور", True, ["شماره سند", "شماره سفارش", "invoice no", "invoice number"]),
    ("invoice_date", "تاریخ فاکتور", False, ["تاریخ", "تاریخ سند", "invoice date"]),
    ("customer_code", "کد مشتری", False, ["کد طرف حساب", "customer code"]),
    ("customer_name", "نام مشتری", True, ["طرف حساب", "نام طرف حساب", "نام خریدار", "customer name"]),
    ("customer_phone", "موبایل", True, ["شماره موبایل", "تلفن همراه", "تلفن", "شماره تماس", "customer phone"]),
    ("customer_email", "ایمیل", False, ["پست الکترونیک", "customer email"]),
    ("product_sku", "کد کالا", True, ["کد محصول", "شناسه کالا", "sku", "product sku"]),
    ("product_name", "شرح کالا", False, ["نام کالا", "نام محصول", "شرح محصول", "product name"]),
    ("unit", "واحد", False, ["واحد کالا"]),
    ("quantity", "تعداد", True, ["مقدار", "quantity", "qty"]),
    ("unit_price", "قیمت واحد", False, ["فی", "نرخ", "مبلغ واحد", "unit price"]),
    ("discount", "تخفیف", False, ["مبلغ تخفیف"]),
    ("tax", "مالیات", False, ["مالیات و عوارض", "مبلغ مالیات"]),
    ("line_total", "مبلغ ردیف", False, ["مبلغ کل", "مبلغ", "جمع ردیف", "line total"]),
    ("currency", "واحد پول", False, ["currency"]),
]


def normalize_header(value):
    text = str(value or "").casefold().replace("ي", "ی").replace("ك", "ک")
    return re.sub(r"[\s_\-\u200c:()/]+", "", text)


def suggest_mapping(headers, fields=None):
    result = {}
    for key, label, _, aliases in fields or FIELDS:
        names = {normalize_header(name) for name in [key, label, *aliases]}
        matches = [index for index, header in enumerate(headers) if normalize_header(header) in names]
        # Repeated names need an explicit choice rather than silently taking one.
        if len(matches) == 1:
            result[key] = matches[0]
    return result


def cell_value(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def read_accounting_export(stream, mapping=None, header_row=None, fields=None):
    field_definitions = fields or FIELDS
    workbook = load_workbook(stream, read_only=True, data_only=True)
    try:
        iterator = workbook.active.iter_rows(values_only=True)
        leading = []
        for _ in range(15):
            cells = next(iterator, None)
            if cells is None:
                break
            leading.append(cells)
        if not leading:
            raise ValueError("فایل خالی است.")
        if header_row is None:
            index = max(range(len(leading)), key=lambda i: len(suggest_mapping(leading[i], field_definitions)))
        else:
            if isinstance(header_row, bool) or not isinstance(header_row, int) or not 1 <= header_row <= len(leading):
                raise ValueError("ردیف عنوان ستون‌ها معتبر نیست؛ یکی از ۱۵ ردیف اول را انتخاب کنید.")
            index = header_row - 1
        headers = [str(value or "").strip() for value in leading[index]]
        detected = suggest_mapping(headers, field_definitions) if mapping is None else mapping
        if not isinstance(detected, dict):
            raise ValueError("انتخاب ستون‌ها معتبر نیست.")
        fields = {field[0] for field in field_definitions}
        for key, col in detected.items():
            if key not in fields or isinstance(col, bool) or not isinstance(col, int) or not 0 <= col < len(headers):
                raise ValueError("انتخاب ستون‌ها معتبر نیست.")
        if len(set(detected.values())) != len(detected):
            raise ValueError("هر ستون فایل را فقط به یک فیلد اختصاص دهید.")
        details = {
            "headers": headers, "mapping": detected, "header_row": index + 1,
            "fields": [{"key": key, "label": label, "required": required} for key, label, required, _ in field_definitions],
            "header_candidates": [{"row": i + 1, "values": [str(v or "") for v in row]} for i, row in enumerate(leading)],
        }
        if any(required and key not in detected for key, _, required, _ in field_definitions):
            return {**details, "mapping_required": True}
        rows = []

        def append_row(number, cells):
            if all(value in (None, "") for value in cells):
                return
            row = {key: cell_value(cells[col] if col < len(cells) else None) for key, col in detected.items()}
            row["row_number"] = number
            rows.append(row)
            if len(rows) > 500:
                raise ValueError("حداکثر ۵۰۰ ردیف در هر فایل دمو مجاز است.")

        for i, cells in enumerate(leading[index + 1:], start=index + 2):
            append_row(i, cells)
        for i, cells in enumerate(iterator, start=len(leading) + 1):
            append_row(i, cells)
        if not rows:
            raise ValueError("فایل هیچ ردیف فاکتوری ندارد.")
        return {**details, "mapping_required": False, "rows": rows}
    finally:
        workbook.close()
