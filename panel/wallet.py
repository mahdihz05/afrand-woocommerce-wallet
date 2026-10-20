"""Wallet-only demo. WordPress owns users, balances and credit idempotency."""
import hashlib
import io
import json
import re
import uuid
from collections import Counter
from decimal import Decimal, InvalidOperation, ROUND_DOWN
from pathlib import Path

from flask import jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .accounting import read_accounting_export
from .domain import normalize_phone
from .woo import WooError

ROOT = 'afrand-wallet'
DIGITS = str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩', '01234567890123456789')
FIELDS = [
    ('reference', 'شماره سند', True, ['شماره فاکتور', 'شماره رسید', 'کد پیگیری', 'reference', 'order_ref', 'invoice no']),
    ('date', 'تاریخ', False, ['تاریخ سند', 'تاریخ فاکتور', 'date']),
    ('customer_name', 'نام مشتری', True, ['نام طرف حساب', 'طرف حساب', 'نام و نام خانوادگی', 'customer name']),
    ('phone', 'موبایل', True, ['شماره موبایل', 'تلفن همراه', 'شماره تماس', 'تلفن', 'customer phone', 'customer_phone']),
    ('amount', 'مبلغ خرید', True, ['مبلغ واریز', 'مبلغ', 'مبلغ کل', 'جمع کل', 'بستانکار', 'مبلغ ردیف', 'amount', 'line_total']),
    ('currency', 'واحد پول', False, ['currency']),
    ('note', 'شرح', False, ['توضیحات', 'شرح سند', 'note']),
]


def amount_value(value, decimals=0):
    if isinstance(value, bool):
        raise ValueError('مبلغ باید عددی بزرگ‌تر از صفر باشد.')
    text = str(value).translate(DIGITS).strip().replace('٬', ',').replace('٫', '.')
    if ',' in text:
        if not re.fullmatch(r'\d{1,3}(?:,\d{3})+(?:\.\d+)?', text):
            raise ValueError('جداکنندهٔ مبلغ معتبر نیست؛ مثلاً ۱۰۰٬۰۰۰ وارد کنید.')
        text = text.replace(',', '')
    if not re.fullmatch(r'\d+(?:\.\d+)?', text):
        raise ValueError('مبلغ باید عددی بزرگ‌تر از صفر باشد.')
    try:
        number = Decimal(text)
        if not number.is_finite() or not 0 < number <= 100000000 or number != number.quantize(Decimal(1).scaleb(-decimals)):
            raise ValueError('مبلغ باید مثبت، حداکثر ۱۰۰ میلیون و با دقت واحد پول کیف پول باشد.')
        return format(number, f'.{decimals}f')
    except (InvalidOperation, TypeError):
        raise ValueError('مبلغ معتبر نیست.') from None


def cashback_values(purchase, percent='5'):
    purchase = amount_value(purchase)
    text = str(percent).translate(DIGITS).strip().replace('٫', '.')
    if not re.fullmatch(r'\d{1,3}(?:\.\d{1,2})?', text) or not 0 < Decimal(text) <= 100:
        raise ValueError('درصد کش‌بک باید بیشتر از صفر و حداکثر ۱۰۰، با حداکثر دو رقم اعشار باشد.')
    rate = Decimal(text)
    amount = (Decimal(purchase) * rate / 100).quantize(Decimal('1'), rounding=ROUND_DOWN)
    if amount < 1:
        raise ValueError('کش‌بک این خرید کمتر از یک واحد پول است؛ مبلغ خرید یا درصد را افزایش دهید.')
    return {'purchase_amount': purchase, 'cashback_percent': format(rate.normalize(), 'f'), 'amount': str(amount)}


def cashback_note(values, note):
    if len(note) > 180:
        raise ValueError('شرح خرید حداکثر ۱۸۰ نویسه است.')
    return f"کش‌بک {values['cashback_percent']}٪ خرید {values['purchase_amount']} {values['currency']}" + (f' — {note}' if note else '')


def customer_values(data):
    if not isinstance(data, dict):
        raise ValueError('اطلاعات مشتری معتبر نیست.')
    phone = normalize_phone(data.get('phone'))
    first = str(data.get('first_name') or '').strip()
    last = str(data.get('last_name') or '').strip()
    if not re.fullmatch(r'09\d{9}', phone) or not first or len(first) > 80 or len(last) > 80:
        raise ValueError('نام و شمارهٔ موبایل معتبر مشتری را وارد کنید.')
    return {'phone': phone, 'first_name': first, 'last_name': last}


def configure_wallet(app, woo, ledger):
    app.extensions['wallet_mode'] = True

    def fail(message, status=400, state=None):
        return jsonify(error=message, category='wallet', **({'state': state} if state else {})), status

    def allowed(write=False):
        if not woo.configured or not woo.host_matches_staging:
            return fail('اتصال وردپرس تنظیم نشده است.', 503)
        if write and (str(app.config.get('WC_WRITES_ENABLED')).lower() not in {'true', '1'} or app.config.get('WC_ENVIRONMENT') != 'staging'):
            return fail('ثبت در کیف پول غیرفعال است.', 403)
        return None

    @app.errorhandler(WooError)
    def wallet_error(error):
        status = error.status if error.status in {400, 401, 403, 404, 409, 503} else 502
        if error.category == 'connection':
            return fail('پاسخ کامل وردپرس دریافت نشد. نتیجهٔ واریز قبلی را بررسی کنید؛ واریز تازه تکرار نکنید.', 502, 'unknown')
        return fail(str(error), status, 'unknown' if status == 409 else None)

    def remote(method, path, **kwargs):
        return woo.request(method, f'{ROOT}/{path}', **kwargs)

    def details():
        data = remote('GET', 'status')
        if not isinstance(data, dict) or not data.get('connected') or data.get('wallet') != 'TeraWallet':
            raise WooError('پاسخ اتصال کیف پول معتبر نیست.', category='api')
        return data

    @app.get('/')
    def index():
        return render_template('wallet.html')

    @app.get('/api/status')
    def status():
        blocked = allowed()
        if blocked:
            return blocked
        data = details()
        return jsonify(**data, writes_enabled=allowed(True) is None, environment='staging')

    @app.get('/api/customers')
    def customers():
        blocked = allowed()
        if blocked:
            return blocked
        try:
            page = int(request.args.get('page', 1))
        except ValueError:
            return fail('صفحهٔ جست‌وجو معتبر نیست.')
        q = request.args.get('q', '').strip()[:100]
        if len(q) < 2 or not 1 <= page <= 100:
            return fail('برای جست‌وجو نام یا موبایل مشتری را وارد کنید.')
        return jsonify(remote('GET', 'customers', params={'q': q, 'page': page}))

    @app.get('/api/customers/<int:customer_id>')
    def customer(customer_id):
        blocked = allowed()
        return blocked or jsonify(remote('GET', f'customers/{customer_id}'))

    @app.post('/api/customers')
    def create_customer():
        blocked = allowed(True)
        if blocked:
            return blocked
        data = request.get_json(silent=True)
        try:
            values = customer_values(data)
        except ValueError as error:
            return fail(str(error))
        if data.get('confirm') is not True:
            return fail('تأیید ثبت مشتری لازم است.')
        return jsonify(remote('POST', 'customers/create', payload={**values, 'confirm': True}))

    def credit_payload(data):
        if not isinstance(data, dict) or data.get('confirm') is not True:
            raise ValueError('تأیید واریز لازم است.')
        request_id = str(data.get('request_id') or '')
        if not re.fullmatch('[a-f0-9]{32,64}', request_id):
            raise ValueError('شناسهٔ درخواست معتبر نیست.')
        customer_id = data.get('customer_id')
        if isinstance(customer_id, bool) or not isinstance(customer_id, int) or customer_id <= 0:
            raise ValueError('مشتری را انتخاب کنید.')
        currency = str(data.get('currency') or '')
        # The currency is fixed by WordPress; the bridge validates it again at write time.
        if currency not in {'IRT', 'IRR'}:
            raise ValueError('واحد پول این دمو باید تومان یا ریال باشد.')
        note = str(data.get('note') or '').strip()
        values = cashback_values(data.get('purchase_amount'), data.get('cashback_percent', '5'))
        note = cashback_note({**values, 'currency': currency}, note)
        key = hashlib.sha256((woo.base_url + ':manual-wallet:' + request_id).encode()).hexdigest()
        return {'request_id': key, 'customer_id': customer_id, 'amount': values['amount'], 'currency': currency, 'note': note, 'confirm': True}, values

    @app.post('/api/wallet/credits')
    def credit():
        blocked = allowed(True)
        if blocked:
            return blocked
        try:
            payload, values = credit_payload(request.get_json(silent=True))
        except ValueError as error:
            return fail(str(error), state='rejected')
        # WordPress owns the durable claim; repeating this key cannot mint new credit.
        result = remote('POST', 'credits', payload=payload)
        return jsonify({**result, **values})

    @app.get('/api/wallet/operations/<request_id>')
    def operation(request_id):
        blocked = allowed()
        if blocked:
            return blocked
        if not re.fullmatch('[a-f0-9]{32,64}', request_id):
            return fail('شناسهٔ درخواست معتبر نیست.')
        key = hashlib.sha256((woo.base_url + ':manual-wallet:' + request_id).encode()).hexdigest()
        return jsonify(remote('GET', f'operations/{key}'))

    def template_bytes():
        workbook = Workbook(); sheet = workbook.active
        sheet.title = 'واریز کیف پول'; sheet.sheet_view.rightToLeft = True
        sheet.append([field[1] for field in FIELDS])
        for cell in sheet[1]:
            cell.font = Font(name='Tahoma', bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='235BAF'); cell.alignment = Alignment(horizontal='right')
        for column, width in {'A':24,'B':18,'C':28,'D':20,'E':20,'F':16,'G':38}.items(): sheet.column_dimensions[column].width = width
        sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = 'A1:G1'
        stream = io.BytesIO(); workbook.save(stream); stream.seek(0); return stream

    @app.get('/template.xlsx')
    def template():
        return send_file(template_bytes(), download_name='afrand-wallet-template.xlsx', as_attachment=True)

    @app.get('/demo.xlsx')
    def sample():
        return send_file(Path(app.root_path).parent / 'demo' / 'afrand-wallet-sample.xlsx', as_attachment=True)

    @app.post('/api/import/preview')
    def preview():
        blocked = allowed()
        if blocked:
            return blocked
        file = request.files.get('file')
        if not file or not file.filename.lower().endswith('.xlsx'):
            return fail('فایل اکسل با پسوند xlsx انتخاب کنید.')
        try:
            mapping = json.loads(request.form['mapping']) if request.form.get('mapping') else None
            header = int(request.form['header_row']) if request.form.get('header_row') else None
            parsed = read_accounting_export(file.stream, mapping, header, fields=FIELDS)
        except Exception:
            return fail('فایل یا انتخاب ستون‌ها معتبر نیست؛ فایل xlsx سالم و حداکثر ۵۰۰ ردیف انتخاب کنید.')
        if parsed['mapping_required']:
            return jsonify(parsed)
        try:
            percent = cashback_values('100000000', request.form.get('cashback_percent', '5'))['cashback_percent']
        except ValueError as error:
            return fail(str(error))
        settings = details(); currency = settings['currency']
        if currency not in {'IRT', 'IRR'}:
            return fail('واحد پول فروشگاه با دموی واریز هماهنگ نیست.')
        rows = parsed.pop('rows'); refs = Counter(str(row.get('reference') or '').strip() for row in rows)
        normalized = []
        for source in rows:
            row = {**source, 'errors': []}
            row['reference'] = str(source.get('reference') or '').strip()
            row['customer_name'] = str(source.get('customer_name') or '').strip()
            row['phone'] = normalize_phone(source.get('phone'))
            row['note'] = str(source.get('note') or '').strip()
            unit = str(source.get('currency') or '').strip().casefold()
            aliases = {'irt':'IRT','تومان':'IRT','تومن':'IRT','irr':'IRR','ریال':'IRR'}
            if unit and aliases.get(unit) != currency: row['errors'].append('واحد پول فایل با کیف پول یکسان نیست؛ تبدیل خودکار انجام نمی‌شود')
            row['currency'] = currency
            if not row['reference'] or len(row['reference']) > 100: row['errors'].append('شماره سند معتبر لازم است')
            if refs[row['reference']] > 1: row['errors'].append('شماره سند در فایل تکراری است؛ هر واریز یک سند یکتا دارد')
            if not row['customer_name'] or len(row['customer_name']) > 80: row['errors'].append('نام مشتری معتبر لازم است')
            if not re.fullmatch(r'09\d{9}', row['phone']): row['errors'].append('موبایل معتبر لازم است؛ صفر ابتدای شماره حفظ شود')
            if len(row['note']) > 180: row['errors'].append('شرح حداکثر ۱۸۰ نویسه است')
            try: row.update(cashback_values(source.get('amount'), percent))
            except ValueError as error: row['errors'].append(str(error))
            normalized.append(row)
        phones = list({row['phone'] for row in normalized if not row['errors']})
        matches = remote('POST', 'customers/lookup', payload={'phones': phones})['matches'] if phones else {}
        for row in normalized:
            if not row['errors']:
                found = matches.get(row['phone'], [])
                if len(found) > 1: row['errors'].append('چند مشتری با این موبایل وجود دارد؛ واریز خودکار انجام نمی‌شود')
                row['customer_id'] = found[0]['id'] if len(found) == 1 else None
                row['customer_action'] = 'مشتری موجود' if row['customer_id'] else 'ساخت مشتری جدید'
        valid = [row for row in normalized if not row['errors']]
        invalid = [row for row in normalized if row['errors']]
        draft = uuid.uuid4().hex
        # A wallet-specific envelope cannot be mistaken for historical order drafts.
        ledger.save_draft(draft, [{'kind':'wallet-cashback-v2', 'rows':valid, 'currency':currency}])
        return jsonify(**parsed, draft_id=draft, valid_rows=valid, invalid_rows=invalid, valid_count=len(valid), invalid_count=len(invalid), currency=currency, cashback_percent=percent, purchase_total=str(sum(Decimal(row['purchase_amount']) for row in valid)), total=str(sum(Decimal(row['amount']) for row in valid)))

    @app.post('/api/import/commit')
    def commit():
        blocked = allowed(True)
        if blocked:
            return blocked
        data = request.get_json(silent=True)
        if not isinstance(data, dict) or data.get('confirm') is not True:
            return fail('تأیید واریز فاکتورها لازم است.')
        draft = ledger.get_draft(str(data.get('draft_id') or ''))
        if not draft or len(draft) != 1 or draft[0].get('kind') != 'wallet-cashback-v2':
            return fail('پیش‌نمایش منقضی یا نامعتبر است؛ فایل را دوباره بررسی کنید.')
        offset = data.get('offset', 0)
        if isinstance(offset, bool) or not isinstance(offset, int) or not 0 <= offset <= len(draft[0]['rows']):
            return fail('مرحلهٔ واریز معتبر نیست.')
        # Bound each request while the UI advances through a larger workbook.
        batch = draft[0]['rows'][offset:offset+2]
        results = []
        for row in batch:
            key = hashlib.sha256((woo.base_url + ':wallet-reference:' + row['reference']).encode()).hexdigest()
            try:
                # Re-match at commit: a preview must never grant stale customer identity.
                found = remote('GET','customers',params={'q':row['phone']})['customers']
                if len(found) > 1: raise WooError('موبایل به چند حساب تعلق دارد.',409)
                customer_id = found[0]['id'] if found else None
                if not customer_id:
                    name = row['customer_name'].split(' ',1)
                    created = remote('POST','customers/create',payload={'first_name':name[0],'last_name':name[1] if len(name)>1 else '', 'phone':row['phone'],'confirm':True})
                    customer_id = created['customer']['id']
                result = remote('POST','credits',payload={'request_id':key,'customer_id':customer_id,'amount':row['amount'],'currency':row['currency'],'note':cashback_note(row, row['note'] or ('سند '+row['reference'])), 'confirm':True})
                results.append({**result,'reference':row['reference'], 'purchase_amount':row['purchase_amount'], 'cashback_percent':row['cashback_percent']})
            except WooError as error:
                results.append({'reference':row['reference'],'state':'needs_review','error':str(error) if error.status else 'پاسخ نامشخص؛ نتیجهٔ سند در وردپرس بررسی شود.'})
        next_offset = offset + len(batch)
        return jsonify(results=results, created=sum(not row.get('replayed') for row in results if row.get('transaction_id')), replayed=sum(bool(row.get('replayed')) for row in results), needs_review=sum(row.get('state')=='needs_review' for row in results), next_offset=next_offset, done=next_offset>=len(draft[0]['rows']))
