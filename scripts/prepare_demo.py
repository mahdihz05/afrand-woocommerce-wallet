"""Prepare marked synthetic catalog fixtures on the configured staging store."""

import io
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from run import app
from panel.woo import WooError
from demo_access import local_test_client


def main():
    client = local_test_client(app)
    status = client.get('/api/status').get_json()
    if not status['connected'] or not status['writes_enabled']:
        raise SystemExit('Staging connection and explicit writes must be enabled first.')
    woo = app.extensions['woo']
    fixtures = []
    for sku, name, price in [
        ('AFRAND-DEMO-001', 'دفتر یادداشت — آزمایشی افرند', '120000'),
        ('AFRAND-DEMO-002', 'خودکار — آزمایشی افرند', '45000'),
    ]:
        found = woo.get_products_by_sku(sku)
        if found:
            product = found[0]
            if not any(m.get('key') == 'afrand_demo_fixture' and m.get('value') == 'synthetic-v1' for m in product.get('meta_data', [])):
                raise SystemExit('A demo SKU belongs to an unmarked product; refusing to reuse it.')
        else:
            product = woo.request('POST', 'products', payload={
                'name': name, 'sku': sku, 'type': 'simple', 'status': 'publish',
                'regular_price': price, 'virtual': True, 'tax_status': 'none',
                'description': 'دادهٔ ساختگی برای نمایش پنل افرند؛ برای فروش واقعی نیست.',
                'meta_data': [{'key': 'afrand_demo_fixture', 'value': 'synthetic-v1'}],
            })
        fixtures.append({'id': product['id'], 'sku': sku})

    from openpyxl import load_workbook
    template = client.get('/template.xlsx')
    workbook = load_workbook(io.BytesIO(template.data))
    sheet = workbook.active
    sheet.delete_rows(2, sheet.max_row)
    for row in [
        ['AFRAND-DEMO-20261006-A', 'مریم آزمایشی افرند', '09120000001', 'afrand.demo.maryam@example.test', 'AFRAND-DEMO-001', 1],
        ['AFRAND-DEMO-20261006-A', 'مریم آزمایشی افرند', '09120000001', 'afrand.demo.maryam@example.test', 'AFRAND-DEMO-002', 2],
        ['AFRAND-DEMO-20261006-B', 'علی آزمایشی افرند', '09120000002', 'afrand.demo.ali@example.test', 'AFRAND-DEMO-001', 1],
        ['AFRAND-DEMO-20261006-C', 'خطای محصول آزمایشی', '09120000003', 'afrand.demo.invalid@example.test', 'SKU-NOT-FOUND', 1],
        ['AFRAND-DEMO-20261006-D', 'خطای تعداد آزمایشی', '09120000004', 'afrand.demo.quantity@example.test', 'AFRAND-DEMO-001', 1.5],
        ['AFRAND-DEMO-20261006-E', 'خطای موبایل آزمایشی', '123', 'afrand.demo.phone@example.test', 'AFRAND-DEMO-001', 1],
    ]:
        sheet.append(row)
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 30
    for row in sheet.iter_rows(min_row=2):
        row[2].number_format = '@'
    target = Path('demo')
    target.mkdir(exist_ok=True)
    workbook.save(target / 'afrand-demo.xlsx')
    Path(app.instance_path, 'demo-fixtures.json').write_text(json.dumps(fixtures, indent=2), encoding='utf-8')
    print('Prepared staging demo products:', fixtures)
    print('Sample workbook: demo/afrand-demo.xlsx (3 valid lines, 2 orders, 3 invalid lines)')


if __name__ == '__main__':
    try:
        main()
    except WooError:
        raise SystemExit('WooCommerce operation failed. Check staging before retrying; credentials are not printed.') from None
