import json
import os
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor

from panel.app import create_app
from panel.domain import normalize_phone
from panel.woo import WooError


class FakeWoo:
    def __init__(self):
        self.created_customers = []
        self.customer_create_calls = 0
        self.customer_create_error = None
        self.customer_create_side_effect_before_error = False
        self.customer_email_search_barrier = None
        self.updated_customers = []
        self.created_orders = []
        self.products_error = None
        self.phone_search_complete = True
        self.customers = [{"id": 7, "first_name": "مریم", "last_name": "نمونه", "phone": "09121234567"}]

    def check_connection(self):
        if getattr(self, "connection_error", None):
            raise RuntimeError(self.connection_error)
        return True

    def get_products(self, search=""):
        if self.products_error:
            raise self.products_error
        return [{"id": 101, "name": "محصول نمونه", "sku": "SKU-1", "price": "120000"}]

    def get_products_by_sku(self, sku):
        return [p for p in self.get_products() if p["sku"] == sku]

    def find_order_by_request_id(self, request_id):
        return next((o for o in self.created_orders if any(m["key"] == "afrand_demo_request_id" and m["value"] == request_id for m in o.get("meta_data", []))), None)

    def search_customers(self, query, page=1, per_page=20):
        barrier = self.customer_email_search_barrier
        if barrier and "@" in query:
            barrier.wait(timeout=5)
        term = query.casefold()
        matches = [c for c in self.customers if term in " ".join(str(c.get(k, "")) for k in ("first_name", "last_name", "email", "phone")).casefold()]
        start = (page - 1) * per_page
        return matches[start:start + per_page], len(matches) > start + per_page

    def get_customer(self, customer_id):
        return next((c for c in self.customers if c["id"] == customer_id), {"id": customer_id})

    def get_customer_orders(self, customer_id):
        return [{"id": 800 + customer_id, "status": "completed", "total": "120000"}]

    def get_product(self, product_id):
        return {"id": product_id, "name": "محصول نمونه", "status": "publish"}

    def search_customers_by_phone(self, phone):
        wanted = normalize_phone(phone)
        return [c for c in self.customers if normalize_phone(c.get("phone")) == wanted], self.phone_search_complete

    def create_customer(self, customer):
        self.customer_create_calls += 1
        if self.customer_create_error and not self.customer_create_side_effect_before_error:
            raise self.customer_create_error
        saved = {"id": 20 + len(self.created_customers), **customer}
        self.created_customers.append(saved)
        self.customers.append(saved)
        if self.customer_create_error:
            raise self.customer_create_error
        return saved

    def update_customer(self, customer_id, customer):
        saved = {**self.get_customer(customer_id), **customer, "id": customer_id}
        self.updated_customers.append(saved)
        for index, current in enumerate(self.customers):
            if current["id"] == customer_id:
                self.customers[index] = saved
                break
        return saved

    def create_order(self, order):
        saved = {"id": 900 + len(self.created_orders), **order}
        self.created_orders.append(saved)
        return saved


class PanelRouteTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.db_path = os.path.join(self.temp.name, "test.sqlite3")
        self.woo = FakeWoo()
        self.base = {
            "TESTING": True,
            "DATABASE": self.db_path,
            "WC_BASE_URL": "https://www.wprokad.ir",
            "WC_STAGING_HOST": "www.wprokad.ir",
            "WC_CONSUMER_KEY": "never-return-this-key",
            "WC_CONSUMER_SECRET": "never-return-this-secret",
            "WC_WRITES_ENABLED": False,
        }
        self.app = create_app(self.base, woo_client=self.woo)
        self.client = self.app.test_client()

    def tearDown(self):
        self.temp.cleanup()

    def test_fractional_quantity_is_rejected_without_writes(self):
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        response = app.test_client().post('/api/orders', json={
            'request_id': 'fractional-order-001', 'confirm': True, 'status': 'on-hold',
            'customer': self.woo.customers[0], 'items': [{'product_id': 101, 'quantity': 1.5}],
        })
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.woo.created_orders, [])

    def test_cross_site_write_is_rejected(self):
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        response = app.test_client().post('/api/orders', json={}, headers={'Origin': 'https://other.example'})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.woo.created_orders, [])

    def test_rejected_product_can_be_corrected_without_creating_customer_first(self):
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        original = self.woo.get_product
        self.woo.get_product = lambda product_id: {**original(product_id), 'stock_status': 'outofstock'}
        payload = {'request_id': 'correctable-order-001', 'confirm': True, 'status': 'on-hold',
                   'customer': {'first_name': 'آزمایشی', 'phone': '09120001111', 'email': 'new@example.test'},
                   'items': [{'product_id': 101, 'quantity': 1}]}
        rejected = app.test_client().post('/api/orders', json=payload)
        self.assertEqual(rejected.get_json()['state'], 'rejected')
        self.assertEqual(self.woo.created_customers, [])
        self.woo.get_product = original
        accepted = app.test_client().post('/api/orders', json=payload)
        self.assertEqual(accepted.status_code, 201)
        self.assertEqual(len(self.woo.created_orders), 1)

    def test_uncertain_order_can_be_reconciled_without_second_write(self):
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        original = self.woo.create_order
        def lost_response(order):
            original(order)
            raise WooError('lost response', category='connection')
        self.woo.create_order = lost_response
        payload = {'request_id': 'lost-order-response-001', 'confirm': True, 'status': 'on-hold',
                   'customer': self.woo.customers[0], 'items': [{'product_id': 101, 'quantity': 1}]}
        first = app.test_client().post('/api/orders', json=payload)
        self.assertEqual(first.status_code, 502)
        replay = app.test_client().post('/api/orders', json=payload)
        self.assertEqual(replay.status_code, 200)
        self.assertTrue(replay.get_json()['replayed'])
        self.assertEqual(len(self.woo.created_orders), 1)

    def test_draft_survives_app_restart_and_commit_replays(self):
        import io
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(['order_ref', 'customer_name', 'customer_phone', 'customer_email', 'product_sku', 'quantity'])
        workbook.active.append(['PERSIST-001', 'مریم', '09121234567', '', 'SKU-1', 1])
        content = io.BytesIO()
        workbook.save(content)
        content.seek(0)
        preview = self.client.post('/api/import/preview', data={'file': (content, 'test.xlsx')}).get_json()
        self.assertEqual(preview['order_count'], 1)
        restarted = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo).test_client()
        payload = {'draft_id': preview['draft_id'], 'status': 'on-hold', 'confirm': True}
        self.assertEqual(restarted.post('/api/import/commit', json=payload).get_json()['created'], 1)
        self.assertEqual(restarted.post('/api/import/commit', json=payload).get_json()['replayed'], 1)
        self.assertEqual(len(self.woo.created_orders), 1)

    def test_customer_can_be_updated_without_creating_a_new_record(self):
        self.woo.customers[0]["email"] = "maryam@example.test"
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        response = app.test_client().put("/api/customers/7", json={
            "first_name": "مریم", "last_name": "به‌روز", "phone": "09121234567", "email": "maryam@example.test", "confirm": True
        })
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(response.get_json()["customer"]["id"], 7)
        self.assertEqual(len(self.woo.updated_customers), 1)
        self.assertEqual(self.woo.created_customers, [])

    def test_customer_update_refuses_email_owned_by_another_customer(self):
        self.woo.customers.extend([
            {"id": 8, "first_name": "مشتری", "phone": "09127778899", "email": "taken@example.test"}
        ])
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        response = app.test_client().put("/api/customers/7", json={
            "first_name": "مریم", "phone": "09121234567", "email": "taken@example.test", "confirm": True
        })
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()["category"], "validation")
        self.assertEqual(self.woo.updated_customers, [])

    def test_customer_search_supports_name_email_and_phone(self):
        self.woo.customers[0]["email"] = "maryam@example.test"
        for query in ("مریم", "maryam@example.test", "09121234567"):
            with self.subTest(query=query):
                response = self.client.get("/api/customers", query_string={"q": query})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.get_json()["customers"][0]["id"], 7)
        self.assertEqual(self.woo.created_customers, [])

    def test_customer_search_exposes_and_serves_results_beyond_the_first_page(self):
        self.woo.customers.extend(
            {"id": customer_id, "first_name": "مریم", "last_name": f"مشتری {customer_id}", "phone": f"0912000{customer_id:04d}"}
            for customer_id in range(100, 121)
        )

        first = self.client.get("/api/customers", query_string={"q": "مریم"})
        second = self.client.get("/api/customers", query_string={"q": "مریم", "page": 2})

        self.assertEqual(first.status_code, 200)
        self.assertEqual(first.get_json()["page"], 1)
        self.assertEqual(first.get_json()["per_page"], 20)
        self.assertTrue(first.get_json()["has_more"])
        self.assertEqual(len(first.get_json()["customers"]), 20)
        self.assertEqual(second.status_code, 200)
        self.assertEqual(second.get_json()["page"], 2)
        self.assertFalse(second.get_json()["has_more"])
        self.assertIn(120, [customer["id"] for customer in second.get_json()["customers"]])

    def test_products_route_returns_store_price(self):
        response = self.client.get("/api/products", query_string={"search": "SKU-1"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["products"][0]["price"], "120000")

    def test_customer_creation_stops_when_phone_dedupe_scan_is_incomplete(self):
        self.woo.phone_search_complete = False
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        response = app.test_client().post("/api/customers", json={
            "first_name": "سارا", "phone": "09125556677", "email": "sara@example.test", "confirm": True
        })
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()["category"], "api")
        self.assertEqual(self.woo.created_customers, [])

    def test_order_refuses_ambiguous_phone_matches(self):
        self.woo.customers.append({"id": 8, "first_name": "مریم", "phone": "09121234567"})
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        response = app.test_client().post("/api/orders", json=self.valid_order())
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()["category"], "validation")
        self.assertEqual(self.woo.created_customers, [])
        self.assertEqual(self.woo.created_orders, [])

    def test_customer_detail_and_orders_are_readable_without_writes(self):
        detail = self.client.get("/api/customers/7")
        orders = self.client.get("/api/customers/7/orders")

        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.get_json()["customer"]["id"], 7)
        self.assertEqual(orders.status_code, 200)
        self.assertEqual(orders.get_json()["orders"][0]["id"], 807)
        self.assertEqual(self.woo.created_customers, [])
        self.assertEqual(self.woo.created_orders, [])

    def test_customer_write_requires_explicit_confirmation(self):
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        response = app.test_client().post("/api/customers", json={
            "first_name": "سارا", "phone": "09125556677", "email": "sara@example.test"
        })
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["category"], "validation")
        self.assertEqual(self.woo.created_customers, [])

    def test_customer_can_be_created_when_email_and_phone_are_unique(self):
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        response = app.test_client().post("/api/customers", json={
            "first_name": "سارا", "last_name": "نمونه", "phone": "09125556677", "email": "sara@example.test", "confirm": True
        })
        self.assertEqual(response.status_code, 201, response.get_json())
        self.assertEqual(response.get_json()["customer"]["id"], 20)
        self.assertEqual(len(self.woo.created_customers), 1)

    def test_customer_creation_rejects_duplicate_phone_without_writing(self):
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        response = app.test_client().post("/api/customers", json={
            "first_name": "سارا", "phone": "09121234567", "email": "sara@example.test", "confirm": True
        })
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()["category"], "validation")
        self.assertEqual(self.woo.created_customers, [])

    def test_standalone_customer_creation_replays_by_normalized_phone(self):
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        client = app.test_client()
        payload = {"first_name": "سارا", "phone": "09125556677", "email": "sara@example.test", "confirm": True}
        first = client.post("/api/customers", json=payload)
        replay_payload = {**payload, "phone": "+98 912 555 6677"}
        replay = client.post("/api/customers", json=replay_payload)

        self.assertEqual(first.status_code, 201, first.get_json())
        self.assertEqual(replay.status_code, 200, replay.get_json())
        self.assertTrue(replay.get_json()["replayed"])
        self.assertEqual(replay.get_json()["customer"]["id"], first.get_json()["customer"]["id"])
        self.assertEqual(self.woo.customer_create_calls, 1)

    def test_concurrent_standalone_customer_posts_for_same_phone_create_once(self):
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        self.woo.customer_email_search_barrier = threading.Barrier(2)
        phone_barrier = threading.Barrier(2)
        phone_search_calls = {"count": 0}
        phone_search_lock = threading.Lock()
        original_phone_search = self.woo.search_customers_by_phone

        def synchronized_phone_search(phone):
            result = original_phone_search(phone)
            with phone_search_lock:
                wait_for_peer = phone_search_calls["count"] < 2
                phone_search_calls["count"] += 1
            if wait_for_peer:
                phone_barrier.wait(timeout=5)
            return result

        self.woo.search_customers_by_phone = synchronized_phone_search
        payload = {"first_name": "سارا", "phone": "09125556677", "email": "sara@example.test", "confirm": True}

        def post_customer():
            return app.test_client().post("/api/customers", json=payload)

        with ThreadPoolExecutor(max_workers=2) as pool:
            responses = list(pool.map(lambda _: post_customer(), range(2)))

        self.assertEqual(self.woo.customer_create_calls, 1)
        self.assertEqual(len(self.woo.created_customers), 1)
        self.assertIn(201, [response.status_code for response in responses])
        self.assertTrue(all(response.status_code in {200, 201, 409} for response in responses))

    def test_uncertain_standalone_customer_create_is_not_retried_when_unresolved(self):
        self.woo.customer_create_error = WooError("simulated timeout", category="connection")
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        client = app.test_client()
        payload = {"first_name": "سارا", "phone": "09125556677", "email": "sara@example.test", "confirm": True}

        first = client.post("/api/customers", json=payload)
        retry = client.post("/api/customers", json=payload)

        self.assertEqual(first.status_code, 502)
        self.assertEqual(first.get_json()["state"], "unknown")
        self.assertEqual(retry.status_code, 409)
        self.assertEqual(retry.get_json()["state"], "unknown")
        self.assertEqual(self.woo.customer_create_calls, 1)
        self.assertEqual(self.woo.created_customers, [])

    def test_uncertain_standalone_customer_create_reconciles_created_customer(self):
        self.woo.customer_create_error = WooError("simulated timeout after create", category="connection")
        self.woo.customer_create_side_effect_before_error = True
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        client = app.test_client()
        payload = {"first_name": "سارا", "phone": "09125556677", "email": "sara@example.test", "confirm": True}

        first = client.post("/api/customers", json=payload)
        reconciled = client.post("/api/customers", json=payload)

        self.assertEqual(first.status_code, 502)
        self.assertEqual(reconciled.status_code, 200)
        self.assertTrue(reconciled.get_json()["replayed"])
        self.assertEqual(reconciled.get_json()["customer"]["phone"], payload["phone"])
        self.assertEqual(self.woo.customer_create_calls, 1)
        self.assertEqual(len(self.woo.created_customers), 1)

    def test_customer_creation_rejects_duplicate_email_without_writing(self):
        self.woo.customers[0]["email"] = "maryam@example.test"
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        response = app.test_client().post("/api/customers", json={
            "first_name": "سارا", "phone": "09125556677", "email": "MARYAM@example.test", "confirm": True
        })
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()["category"], "validation")
        self.assertEqual(self.woo.created_customers, [])

    def test_demo_page_has_customer_details_and_direct_save_controls(self):
        response = self.client.get("/")
        page = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn('id="customer-details"', page)
        self.assertIn('id="load-more-customers"', page)
        self.assertIn('id="customer-edit-form"', page)
        self.assertIn('id="save-customer-button"', page)

    def test_status_never_returns_api_secrets_and_starts_read_only(self):
        response = self.client.get("/api/status")
        body = json.dumps(response.get_json())
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("never-return-this-key", body)
        self.assertNotIn("never-return-this-secret", body)
        self.assertFalse(response.get_json()["writes_enabled"])
        self.assertEqual(response.get_json()["environment"], "staging")

    def test_api_errors_are_categorized_without_echoing_upstream_details(self):
        app = create_app(self.base, woo_client=self.woo)
        client = app.test_client()
        cases = [(401, "unauthorized"), (403, "forbidden"), (502, "api"), (None, "connection")]
        for code, category in cases:
            with self.subTest(code=code):
                self.woo.products_error = WooError("consumer-secret must never leak", code)
                response = client.get("/api/products")
                self.assertEqual(response.get_json()["category"], category)
                self.assertNotIn("consumer-secret", json.dumps(response.get_json()))
        self.woo.products_error = WooError("consumer-secret must never leak", category="api")
        response = client.get("/api/products")
        self.assertEqual(response.get_json()["category"], "api")
        self.assertNotIn("consumer-secret", json.dumps(response.get_json()))
        self.woo.products_error = None

    def test_malformed_order_customer_is_a_validation_error_not_a_traceback(self):
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        payload = self.valid_order()
        payload["customer"] = "not-an-object"
        response = app.test_client().post("/api/orders", json=payload)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["category"], "validation")

    def test_customer_validation_error_has_validation_category(self):
        response = self.client.post("/api/customers", json={"first_name": "بدون ایمیل"})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.get_json()["category"], "forbidden")
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        response = app.test_client().post("/api/customers", json={"first_name": "بدون ایمیل", "confirm": True})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["category"], "validation")

    def test_status_checks_live_api_health_not_only_config_presence(self):
        self.woo.connection_error = "temporary upstream failure"
        response = self.client.get("/api/status")
        self.assertFalse(response.get_json()["connected"])
        self.assertIn("api_error", response.get_json())
        self.assertEqual(response.get_json()["api_error_category"], "connection")
        self.assertNotIn("never-return-this-secret", json.dumps(response.get_json()))

    def test_status_labels_unapproved_store_host_as_connection_error(self):
        app = create_app({**self.base, "WC_BASE_URL": "https://wrong.example"}, woo_client=self.woo)
        response = app.test_client().get("/api/status")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["api_error_category"], "connection")
        self.assertIn("میزبان", response.get_json()["api_error"])

    def test_order_write_is_blocked_by_default(self):
        response = self.client.post("/api/orders", json=self.valid_order())
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.get_json()["category"], "forbidden")
        self.assertEqual(self.woo.created_orders, [])
        self.assertEqual(self.woo.created_customers, [])

    def test_unknown_store_host_cannot_write_even_when_enabled(self):
        config = {**self.base, "WC_WRITES_ENABLED": True, "WC_BASE_URL": "https://not-the-staging.example"}
        app = create_app(config, woo_client=self.woo)
        response = app.test_client().post("/api/orders", json=self.valid_order())
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.woo.created_orders, [])

    def test_order_requires_explicit_confirmation_and_status(self):
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        payload = self.valid_order()
        payload.pop("status")
        response = app.test_client().post("/api/orders", json=payload)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.woo.created_orders, [])

        payload = self.valid_order()
        payload["confirm"] = False
        response = app.test_client().post("/api/orders", json=payload)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.woo.created_orders, [])

    def test_manual_order_rejects_phone_shorter_than_ten_normalized_digits(self):
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        payload = self.valid_order()
        payload["customer"]["phone"] = "123456789"

        response = app.test_client().post("/api/orders", json=payload)

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["category"], "validation")
        self.assertEqual(self.woo.customer_create_calls, 0)
        self.assertEqual(self.woo.created_orders, [])

    def test_confirmed_order_matches_customer_and_sets_in_person_marker(self):
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        response = app.test_client().post("/api/orders", json=self.valid_order())
        self.assertEqual(response.status_code, 201, response.get_json())
        self.assertEqual(len(self.woo.created_customers), 0)
        order = self.woo.created_orders[0]
        self.assertEqual(order["customer_id"], 7)
        self.assertEqual(order["status"], "pending")
        self.assertIn({"key": "afrand_order_source", "value": "in_person"}, order["meta_data"])

    def test_replaying_same_request_id_does_not_create_a_second_order(self):
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        client = app.test_client()
        first = client.post("/api/orders", json=self.valid_order())
        second = client.post("/api/orders", json=self.valid_order())
        self.assertEqual(first.status_code, 201)
        self.assertEqual(second.status_code, 200)
        self.assertTrue(second.get_json()["replayed"])
        self.assertEqual(len(self.woo.created_orders), 1)

    def test_paid_or_processing_status_is_rejected_until_payment_policy_is_decided(self):
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        payload = self.valid_order()
        payload["status"] = "completed"
        response = app.test_client().post("/api/orders", json=payload)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.woo.created_orders, [])

    def test_new_customer_is_created_before_order_and_requires_email(self):
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        payload = self.valid_order()
        payload["request_id"] = "new-customer-001"
        payload["customer"] = {"phone": "09125556677", "first_name": "سارا", "last_name": "نمونه", "email": "sara@example.test"}
        response = app.test_client().post("/api/orders", json=payload)
        self.assertEqual(response.status_code, 201, response.get_json())
        self.assertEqual(len(self.woo.created_customers), 1)
        self.assertEqual(self.woo.created_orders[0]["customer_id"], self.woo.created_customers[0]["id"])

    def test_order_does_not_create_customer_with_email_owned_by_another_record(self):
        self.woo.customers[0]["email"] = "maryam@example.test"
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        payload = self.valid_order()
        payload["request_id"] = "email-duplicate-001"
        payload["customer"] = {"phone": "09125556677", "first_name": "سارا", "email": "MARYAM@example.test"}

        response = app.test_client().post("/api/orders", json=payload)

        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()["category"], "validation")
        self.assertEqual(self.woo.created_customers, [])
        self.assertEqual(self.woo.created_orders, [])

    def test_order_stops_on_exact_email_match_without_customer_id(self):
        self.woo.customers.append({"first_name": "مشتری", "phone": "09127778899", "email": "unknown-id@example.test"})
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        payload = self.valid_order()
        payload["request_id"] = "email-no-id-001"
        payload["customer"] = {"phone": "09125556677", "first_name": "سارا", "email": "unknown-id@example.test"}
        response = app.test_client().post("/api/orders", json=payload)
        self.assertEqual(response.status_code, 409)
        self.assertEqual(self.woo.created_customers, [])
        self.assertEqual(self.woo.created_orders, [])

    def test_new_customer_without_email_is_not_created(self):
        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        payload = self.valid_order()
        payload["request_id"] = "new-customer-002"
        payload["customer"] = {"phone": "09125556677", "first_name": "سارا", "last_name": "نمونه"}
        response = app.test_client().post("/api/orders", json=payload)
        self.assertEqual(response.status_code, 409)
        self.assertEqual(self.woo.created_customers, [])
        self.assertEqual(self.woo.created_orders, [])

    def test_import_preview_rejects_phone_shorter_than_ten_normalized_digits_before_writes(self):
        import io
        from openpyxl import Workbook

        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        client = app.test_client()
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["order_ref", "customer_name", "customer_phone", "customer_email", "product_sku", "quantity"])
        sheet.append(["short-phone-001", "سارا نمونه", "123456789", "sara@example.test", "SKU-1", 1])
        stream = io.BytesIO()
        workbook.save(stream)
        stream.seek(0)

        preview = client.post("/api/import/preview", data={"file": (stream, "orders.xlsx")}, content_type="multipart/form-data")
        self.assertEqual(preview.status_code, 200, preview.get_json())
        self.assertEqual(preview.get_json()["valid_count"], 0)
        self.assertEqual(preview.get_json()["invalid_count"], 1)
        commit = client.post("/api/import/commit", json={"draft_id": preview.get_json()["draft_id"], "status": "pending", "confirm": True})
        self.assertEqual(commit.status_code, 404)
        self.assertEqual(self.woo.customer_create_calls, 0)
        self.assertEqual(self.woo.created_orders, [])

    def test_import_order_references_with_distinct_persian_text_have_distinct_idempotency_keys(self):
        import io
        from openpyxl import Workbook

        app = create_app({**self.base, "WC_WRITES_ENABLED": True}, woo_client=self.woo)
        client = app.test_client()
        refs = ["سفارش-الف", "سفارش-بیت"]

        def preview():
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["order_ref", "customer_name", "customer_phone", "customer_email", "product_sku", "quantity"])
            for ref in refs:
                sheet.append([ref, "مریم نمونه", "09121234567", "maryam@example.test", "SKU-1", 1])
            stream = io.BytesIO()
            workbook.save(stream)
            stream.seek(0)
            return client.post("/api/import/preview", data={"file": (stream, "orders.xlsx")}, content_type="multipart/form-data")

        draft = preview()
        self.assertEqual(draft.status_code, 200, draft.get_json())
        committed = client.post("/api/import/commit", json={"draft_id": draft.get_json()["draft_id"], "status": "pending", "confirm": True})
        self.assertEqual(committed.status_code, 200, committed.get_json())
        self.assertEqual(committed.get_json()["created"], 2)
        self.assertEqual(len(self.woo.created_orders), 2)
        for order, ref in zip(self.woo.created_orders, refs):
            self.assertIn({"key": "afrand_import_reference", "value": ref}, order["meta_data"])

        replay_draft = preview()
        replayed = client.post("/api/import/commit", json={"draft_id": replay_draft.get_json()["draft_id"], "status": "pending", "confirm": True})
        self.assertEqual(replayed.status_code, 200)
        self.assertEqual(replayed.get_json()["replayed"], 2)
        self.assertEqual(len(self.woo.created_orders), 2)

    def test_import_template_is_downloadable_xlsx(self):
        response = self.client.get("/template.xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response.content_type)
        self.assertTrue(response.data.startswith(b"PK"))
        response.close()

    def test_preview_rejects_non_xlsx_before_any_write(self):
        import io
        response = self.client.post("/api/import/preview", data={"file": (io.BytesIO(b"bad"), "orders.csv")}, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.woo.created_orders, [])

    @staticmethod
    def valid_order():
        return {
            "request_id": "demo-operation-001",
            "confirm": True,
            "customer": {"phone": "09121234567", "first_name": "مریم", "last_name": "نمونه"},
            "items": [{"product_id": 101, "quantity": 1}],
            "status": "pending",
        }


if __name__ == "__main__":
    unittest.main()
