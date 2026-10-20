import io
import json
import tempfile
import threading
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from panel.app import create_app
from panel.domain import normalize_phone
from panel.wallet import amount_value, cashback_values
from panel.woo import WooError


class FakeWallet:
    configured = True
    host_matches_staging = True
    base_url = 'https://www.wprokad.ir'

    def __init__(self):
        self.customers = [{'id':7,'first_name':'مریم','last_name':'نمونه','phone':'09121234567'}]
        self.transactions = []
        self.operations = {}
        self.calls = []
        self.currency = 'IRT'
        self.lose_credit_response = False
        self.lock = threading.Lock()

    def request(self, method, path, params=None, payload=None):
        self.calls.append((method,path))
        if not path.startswith('afrand-wallet/'):
            raise AssertionError('Wallet mode contacted a commerce endpoint: '+path)
        route = path.removeprefix('afrand-wallet/')
        if route == 'status':
            return {'connected':True,'wallet':'TeraWallet','currency':self.currency,'decimals':0}
        if route == 'customers' and method == 'GET':
            q = params['q']; phone = normalize_phone(q)
            found = [c for c in self.customers if (c['phone'] == phone if phone.startswith('09') else q in c['first_name']+' '+c['last_name'])]
            return {'customers':found,'page':1,'has_more':False}
        if route == 'customers/create':
            found = [c for c in self.customers if c['phone'] == normalize_phone(payload['phone'])]
            if len(found) > 1: raise WooError('ambiguous phone',409)
            if found: return {'customer':found[0],'replayed':True}
            c = {k:v for k,v in payload.items() if k in {'first_name','last_name','phone'}}
            c['id'] = 20+len(self.customers); self.customers.append(c)
            return {'customer':c,'replayed':False}
        if route == 'customers/lookup':
            return {'matches':{phone:[c for c in self.customers if c['phone']==phone] for phone in payload['phones']}}
        if route.startswith('customers/'):
            customer = next((c for c in self.customers if c['id'] == int(route.split('/')[1])),None)
            if not customer: raise WooError('not found',404)
            rows = [t for t in self.transactions if t['customer_id'] == customer['id']]
            return {'customer':customer,'wallet':{'balance':str(sum(int(t['amount']) for t in rows)),'currency':self.currency,
                    'transactions':[{**t,'type':'credit','date':'2026-10-08','details':t['note']} for t in rows]}}
        if route.startswith('operations/'):
            key = route.split('/')[1]
            if key not in self.operations: raise WooError('not found',404)
            return {**self.operations[key], 'replayed':True}
        if route == 'credits':
            with self.lock:
                if payload['currency'] != self.currency: raise WooError('wrong currency',400)
                key = payload['request_id']
                if key in self.operations:
                    prior = self.operations[key]
                    if any(prior[k] != payload[k] for k in ('customer_id','amount','currency','note')): raise WooError('request conflict',409)
                    return {**prior,'replayed':True}
                if not any(c['id'] == payload['customer_id'] for c in self.customers): raise WooError('missing user',400)
                saved = {**payload,'transaction_id':100+len(self.transactions),'state':'succeeded','replayed':False}
                self.operations[key] = saved; self.transactions.append(saved)
                if self.lose_credit_response: self.lose_credit_response = False; raise WooError('lost response',category='connection')
                return saved
        raise AssertionError((method,path))


def workbook(rows, headers=None):
    wb = Workbook(); sheet = wb.active
    sheet.append(headers or ['شماره سند','نام مشتری','موبایل','مبلغ واریز','واحد پول'])
    for row in rows: sheet.append(row)
    data = io.BytesIO(); wb.save(data); data.seek(0); return data


class WalletTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.woo = FakeWallet()
        self.config = {'TESTING':True,'WALLET_MODE':True,'DATABASE':str(Path(self.temp.name)/'wallet.db'),
            'WC_BASE_URL':self.woo.base_url,'WC_STAGING_HOST':'www.wprokad.ir','WC_CONSUMER_KEY':'test','WC_CONSUMER_SECRET':'test','WC_WRITES_ENABLED':True}
        self.app = create_app(self.config,woo_client=self.woo); self.client = self.app.test_client()

    def tearDown(self): self.temp.cleanup()

    def preview(self, rows):
        return self.client.post('/api/import/preview',data={'file':(workbook(rows),'wallet.xlsx')})

    def payload(self, **changes):
        return {'request_id':'a'*32,'customer_id':7,'purchase_amount':'25000','currency':'IRT','note':'test','confirm':True,**changes}

    def test_wallet_only_routes_and_no_email_fields(self):
        self.assertEqual(self.client.get('/api/status').json['currency'],'IRT')
        html = self.client.get('/').get_data(as_text=True)
        self.assertNotIn('type="email"',html); self.assertNotIn('product-query',html)
        for method,path in [('GET','/api/products'),('POST','/api/orders')]: self.assertEqual(self.client.open(path,method=method).status_code,404)

    def test_create_no_email_and_normalized_phone_replay(self):
        data={'first_name':'ندا','phone':'۰۹۱۲۳۳۳۴۴۵۵','confirm':True}
        first=self.client.post('/api/customers',json=data); second=self.client.post('/api/customers',json={**data,'phone':'+989123334455'})
        self.assertEqual(first.status_code,200); self.assertEqual(first.json['customer']['id'],second.json['customer']['id'])
        self.assertTrue(second.json['replayed']); self.assertNotIn('email',first.json['customer'])

    def test_customer_requires_name_valid_mobile_confirm(self):
        for data in ({'phone':'09123334455','confirm':True},{'first_name':'ندا','phone':'123','confirm':True},{'first_name':'ندا','phone':'09123334455','confirm':False}):
            self.assertEqual(self.client.post('/api/customers',json=data).status_code,400)
        self.assertEqual(len(self.woo.customers),1)

    def test_credit_idempotency_and_changed_payload_conflict(self):
        first=self.client.post('/api/wallet/credits',json=self.payload())
        again=self.client.post('/api/wallet/credits',json=self.payload())
        conflict=self.client.post('/api/wallet/credits',json=self.payload(purchase_amount='30000'))
        self.assertEqual(first.status_code,200); self.assertTrue(again.json['replayed']); self.assertEqual(conflict.status_code,409)
        self.assertEqual(len(self.woo.transactions),1)

    def test_lost_response_reconciles_without_new_credit(self):
        self.woo.lose_credit_response=True
        self.assertEqual(self.client.post('/api/wallet/credits',json=self.payload()).status_code,502)
        result=self.client.get('/api/wallet/operations/'+'a'*32)
        self.assertEqual(result.status_code,200); self.assertTrue(result.json['replayed']); self.assertEqual(len(self.woo.transactions),1)

    def test_credit_validation_prevents_all_writes(self):
        for changes in ({'purchase_amount':'-1'},{'purchase_amount':True},{'purchase_amount':'NaN'},{'purchase_amount':'100000001'},{'purchase_amount':'1.5'},{'confirm':False},{'customer_id':True},{'request_id':'x'},{'currency':'USD'},{'note':'x'*181}):
            self.assertEqual(self.client.post('/api/wallet/credits',json=self.payload(**changes)).status_code,400,changes)
        self.assertFalse(self.woo.transactions)

    def test_wallet_currency_is_checked_at_write(self):
        self.woo.currency='IRR'
        self.assertEqual(self.client.post('/api/wallet/credits',json=self.payload()).status_code,400)
        self.assertFalse(self.woo.transactions)

    def test_customer_wallet_history(self):
        self.client.post('/api/wallet/credits',json=self.payload())
        data=self.client.get('/api/customers/7').json
        self.assertEqual(data['wallet']['balance'],'1250'); self.assertEqual(data['wallet']['transactions'][0]['transaction_id'],100)

    def test_readonly_and_host_guard(self):
        self.app.config['WC_WRITES_ENABLED']=False
        self.assertEqual(self.client.post('/api/wallet/credits',json=self.payload()).status_code,403)
        self.woo.host_matches_staging=False
        self.assertEqual(self.client.get('/api/status').status_code,503)

    def test_cross_origin_mutation_rejected(self):
        self.assertEqual(self.client.post('/api/wallet/credits',json=self.payload(),headers={'Origin':'https://example.invalid'}).status_code,403)
        self.assertFalse(self.woo.transactions)

    def test_excel_preview_and_commit_existing_new_invalid_repeat(self):
        rows=[['S1','مریم نمونه','09121234567',1000,'تومان'],['S2','ندا نمونه','09123334455',2000,'تومان'],['S3','ناقص','123',-50,'تومان']]
        result=self.preview(rows).json
        self.assertEqual((result['valid_count'],result['invalid_count'],result['total']),(2,1,'150'))
        self.assertEqual(result['purchase_total'],'3000')
        self.assertFalse(self.woo.transactions); self.assertEqual(len(self.woo.customers),1)
        payload={'draft_id':result['draft_id'],'confirm':True}
        committed=self.client.post('/api/import/commit',json=payload).json
        self.assertEqual(committed['created'],2); self.assertEqual(len(self.woo.customers),2)
        # Reupload changed file name / row order cannot mint duplicate deposits.
        new=self.preview(list(reversed(rows))).json
        replay=self.client.post('/api/import/commit',json={'draft_id':new['draft_id'],'confirm':True}).json
        self.assertEqual(replay['replayed'],2); self.assertEqual(replay['created'],0); self.assertEqual(len(self.woo.transactions),2)

    def test_duplicate_reference_and_wrong_currency_invalid(self):
        result=self.preview([['DUP','مریم','09121234567',100,'IRT'],['DUP','مریم','09121234567',100,'IRT'],['IRR','مریم','09121234567',100,'ریال']]).json
        self.assertEqual(result['invalid_count'],3); self.assertEqual(result['valid_count'],0)

    def test_ambiguous_phone_never_credits(self):
        self.woo.customers.append({**self.woo.customers[0],'id':8})
        result=self.preview([['S1','مریم','09121234567',100,'IRT']]).json
        self.assertEqual(result['invalid_count'],1)

    def test_identity_rechecked_after_preview(self):
        result=self.preview([['S1','مریم','09121234567',100,'IRT']]).json
        self.woo.customers.append({**self.woo.customers[0],'id':8})
        response=self.client.post('/api/import/commit',json={'draft_id':result['draft_id'],'confirm':True}).json
        self.assertEqual(response['needs_review'],1); self.assertFalse(self.woo.transactions)

    def test_import_unknown_response_repeat_same_reference(self):
        result=self.preview([['S1','مریم','09121234567',100,'IRT']]).json
        payload={'draft_id':result['draft_id'],'confirm':True}; self.woo.lose_credit_response=True
        first=self.client.post('/api/import/commit',json=payload).json
        self.assertEqual(first['needs_review'],1)
        second=self.client.post('/api/import/commit',json=payload).json
        self.assertEqual(second['replayed'],1); self.assertEqual(len(self.woo.transactions),1)

    def test_changed_import_amount_same_reference_blocked(self):
        for amount in (100,200):
            result=self.preview([['S1','مریم','09121234567',amount,'IRT']]).json
            response=self.client.post('/api/import/commit',json={'draft_id':result['draft_id'],'confirm':True}).json
        self.assertEqual(response['needs_review'],1); self.assertEqual(len(self.woo.transactions),1)

    def test_old_order_draft_cannot_be_committed_as_wallet(self):
        self.app.extensions['ledger'].save_draft('old',[{'order_ref':'old'}])
        self.assertEqual(self.client.post('/api/import/commit',json={'draft_id':'old','confirm':True}).status_code,400)
        self.assertFalse(self.woo.transactions)

    def test_large_import_advances_in_bounded_batches(self):
        result=self.preview([[f'B{i}','مریم','09121234567',100,'IRT'] for i in range(5)]).json
        offset=0; created=0
        while True:
            response=self.client.post('/api/import/commit',json={'draft_id':result['draft_id'],'confirm':True,'offset':offset}).json
            self.assertLessEqual(len(response['results']),2); created+=response['created']
            if response['done']: break
            self.assertGreater(response['next_offset'],offset); offset=response['next_offset']
        self.assertEqual(created,5); self.assertEqual(len(self.woo.transactions),5)

    def test_excel_mapping_extra_columns_heading_and_500_limit(self):
        data=workbook([['junk','Other']],headers=['گزارش'])
        response=self.client.post('/api/import/preview',data={'file':(data,'wallet.xlsx')})
        self.assertTrue(response.json['mapping_required'])
        self.assertNotIn('email',[f['key'] for f in response.json['fields']])
        too_many=self.preview([[f'S{i}','مریم','09121234567',100,'IRT'] for i in range(501)])
        self.assertEqual(too_many.status_code,400)

    def test_template_download_has_wallet_columns_no_email_product(self):
        data=self.client.get('/template.xlsx').data; wb=load_workbook(io.BytesIO(data),read_only=True)
        headers=next(wb.active.values); wb.close()
        self.assertIn('مبلغ خرید',headers); self.assertNotIn('ایمیل',headers); self.assertNotIn('کد کالا',headers)

    def test_server_credits_cashback_only_and_ignores_client_amount(self):
        response=self.client.post('/api/wallet/credits',json=self.payload(purchase_amount='100000',amount='100000'))
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json['amount'],'5000')
        self.assertEqual(self.woo.transactions[0]['amount'],'5000')
        self.assertIn('خرید 100000 IRT',self.woo.transactions[0]['note'])

    def test_custom_percentage_rounding_and_persian_digits(self):
        response=self.client.post('/api/wallet/credits',json=self.payload(purchase_amount='۱۰۰٬۰۰۱',cashback_percent='۷٫۲۵'))
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json['amount'],'7250')
        self.assertEqual(response.json['cashback_percent'],'7.25')
        self.assertEqual(cashback_values('19','100')['amount'],'19')

    def test_invalid_percent_and_tiny_credit_never_write(self):
        for percent in ('0','-1','100.01','1.001','NaN','inf',True,None,'1e2',''):
            self.assertEqual(self.client.post('/api/wallet/credits',json=self.payload(cashback_percent=percent)).status_code,400,percent)
        self.assertEqual(self.client.post('/api/wallet/credits',json=self.payload(purchase_amount='19')).status_code,400)
        self.assertFalse(self.woo.transactions)

    def test_legacy_full_credit_payload_and_draft_blocked(self):
        payload=self.payload(); payload['amount']=payload.pop('purchase_amount')
        self.assertEqual(self.client.post('/api/wallet/credits',json=payload).status_code,400)
        self.app.extensions['ledger'].save_draft('old-wallet',[{'kind':'wallet-v1','rows':[{'amount':'100000'}]}])
        self.assertEqual(self.client.post('/api/import/commit',json={'draft_id':'old-wallet','confirm':True}).status_code,400)
        self.assertFalse(self.woo.transactions)

    def test_same_credit_different_purchase_with_same_rounded_amount_conflicts(self):
        self.client.post('/api/wallet/credits',json=self.payload(purchase_amount='100000'))
        response=self.client.post('/api/wallet/credits',json=self.payload(purchase_amount='100001'))
        self.assertEqual(response.status_code,409)
        self.assertEqual(len(self.woo.transactions),1)

    def test_import_custom_percent_frozen_and_changed_rate_conflict(self):
        rows=[['RATE','مریم','09121234567',100000,'IRT']]
        preview=self.client.post('/api/import/preview',data={'file':(workbook(rows),'rate.xlsx'),'cashback_percent':'7.25'}).json
        row=preview['valid_rows'][0]
        self.assertEqual((row['purchase_amount'],row['amount'],row['cashback_percent']),('100000','7250','7.25'))
        commit=self.client.post('/api/import/commit',json={'draft_id':preview['draft_id'],'cashback_percent':'100','confirm':True}).json
        self.assertEqual(commit['results'][0]['amount'],'7250')
        changed=self.preview(rows).json
        response=self.client.post('/api/import/commit',json={'draft_id':changed['draft_id'],'confirm':True}).json
        self.assertEqual(response['needs_review'],1)
        self.assertEqual(len(self.woo.transactions),1)

    def test_amount_normalization(self):
        self.assertEqual(amount_value('۱۲۳٬۴۵۶'),'123456')
        self.assertEqual(amount_value('٠١٢٣'),'123')
        for value in ('1,23','-1','1e3','inf',False,'0','12.1'):
            with self.assertRaises(ValueError): amount_value(value)


if __name__ == '__main__': unittest.main()
