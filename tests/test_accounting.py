import io
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from panel.accounting import read_accounting_export
from panel.app import create_app
from test_app import FakeWoo


def export(rows):
    # Small in-memory fixtures for the parser; no deliverable workbook authoring.
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


class AccountingTests(unittest.TestCase):
    def test_reordered_persian_columns_extra_columns_and_title_rows(self):
        result = read_accounting_export(export([
            ['گزارش فروش روزانه'], [],
            ['شرح کالا', 'تعداد', 'نام مشتري', 'کد کالا', 'موبايل', 'شماره فاکتور', 'ستون اضافی', 'قیمت واحد'],
            ['دفتر', '۲', 'مریم', 'SKU-1', 9121234567, 1001, 'نادیده', 120000],
        ]))
        self.assertEqual(result['header_row'], 3)
        self.assertFalse(result['mapping_required'])
        self.assertEqual(result['rows'][0]['row_number'], 4)
        self.assertEqual(result['rows'][0]['product_name'], 'دفتر')
        self.assertEqual(result['rows'][0]['order_ref'], 1001)
        self.assertEqual(result['rows'][0]['unit_price'], 120000)

    def test_unknown_titles_require_mapping_without_reading_commerce(self):
        result = read_accounting_export(export([['سند فروش', 'خریدار', 'تماس', 'شناسه', 'مقدار فروش'], ['A', 'مریم', '09121234567', 'SKU-1', 1]]))
        self.assertTrue(result['mapping_required'])
        self.assertNotIn('rows', result)

    def test_explicit_mapping_and_selected_header_row(self):
        result = read_accounting_export(export([['گزارش'], ['X', 'Y', 'Z', 'P', 'Q'], ['A', 'مریم', '09121234567', 'SKU-1', 1]]),
            {'order_ref': 0, 'customer_name': 1, 'customer_phone': 2, 'product_sku': 3, 'quantity': 4}, 2)
        self.assertEqual(result['rows'][0]['order_ref'], 'A')

    def test_rejects_duplicate_column_mapping(self):
        with self.assertRaisesRegex(ValueError, 'یک فیلد'):
            read_accounting_export(export([['A', 'B']]), {'order_ref': 0, 'customer_name': 0})

    def test_rejects_out_of_range_and_boolean_mapping(self):
        for value in [-1, 20, True, '0']:
            with self.subTest(value=value), self.assertRaises(ValueError):
                read_accounting_export(export([['A', 'B']]), {'order_ref': value})

    def test_duplicate_header_needs_manual_choice(self):
        result = read_accounting_export(export([['شماره فاکتور', 'موبایل', 'نام مشتری', 'کد کالا', 'تعداد', 'تعداد']]))
        self.assertTrue(result['mapping_required'])
        self.assertNotIn('quantity', result['mapping'])

    def test_datetime_and_integer_identifier_normalization(self):
        result = read_accounting_export(export([
            ['شماره فاکتور', 'موبایل', 'نام مشتری', 'کد کالا', 'تعداد', 'تاریخ'],
            [1001.0, '09121234567', 'مریم', 'SKU-1', 1, datetime(2026, 10, 6)],
        ]))
        self.assertEqual(result['rows'][0]['order_ref'], 1001)
        self.assertTrue(result['rows'][0]['invoice_date'].startswith('2026-10-06'))

    def test_row_limit(self):
        with self.assertRaisesRegex(ValueError, '۵۰۰'):
            read_accounting_export(export([['شماره فاکتور', 'موبایل', 'نام مشتری', 'کد کالا', 'تعداد']] +
                [[str(i), '09121234567', 'مریم', 'SKU-1', 1] for i in range(501)]))

    def test_old_english_template_remains_accepted(self):
        result = read_accounting_export(export([
            ['order_ref', 'customer_name', 'customer_phone', 'customer_email', 'product_sku', 'quantity'],
            ['OLD', 'مریم', '09121234567', '', 'SKU-1', 1],
        ]))
        self.assertFalse(result['mapping_required'])

    def test_accounting_preview_commit_metadata_and_replay(self):
        with tempfile.TemporaryDirectory() as folder:
            woo = FakeWoo()
            app = create_app({'TESTING': True, 'DATABASE': str(Path(folder) / 'test.db'),
                'WC_BASE_URL': 'https://www.wprokad.ir', 'WC_STAGING_HOST': 'www.wprokad.ir',
                'WC_CONSUMER_KEY': 'test', 'WC_CONSUMER_SECRET': 'test', 'WC_WRITES_ENABLED': True}, woo_client=woo)
            client = app.test_client()
            response = client.post('/api/import/preview', data={'file': (export([
                ['شماره فاکتور', 'نام مشتری', 'موبایل', 'کد کالا', 'تعداد', 'مبلغ ردیف', 'واحد پول'],
                ['ACC-TEST', 'مریم', '09121234567', 'SKU-1', 2, 999999, 'ریال'],
            ]), 'accounting.xlsx')})
            self.assertEqual(response.status_code, 200)
            data = response.get_json()
            self.assertEqual(data['valid_count'], 1)
            self.assertEqual(woo.created_orders, [])
            payload = {'draft_id': data['draft_id'], 'confirm': True, 'status': 'on-hold'}
            self.assertEqual(client.post('/api/import/commit', json=payload).get_json()['created'], 1)
            order = woo.created_orders[0]
            self.assertEqual(order['line_items'], [{'product_id': 101, 'quantity': 2}])
            self.assertNotIn('total', order['line_items'][0])
            metadata = {m['key']: m['value'] for m in order['meta_data']}
            self.assertEqual(metadata['afrand_accounting_export'][0]['line_total'], 999999)
            self.assertEqual(metadata['afrand_accounting_export'][0]['currency'], 'ریال')
            self.assertEqual(client.post('/api/import/commit', json=payload).get_json()['replayed'], 1)
            self.assertEqual(len(woo.created_orders), 1)

    def test_packaged_sample_is_typed_and_cached(self):
        workbook = load_workbook(io.BytesIO(Path('demo/afrand-accounting-sample.xlsx').read_bytes()), data_only=True)
        sheet = workbook.active
        self.assertEqual(sheet['E2'].value, '09120000001')
        self.assertEqual(sheet['N3'].value, 90000)
        self.assertIsInstance(sheet['B2'].value, datetime)
        self.assertEqual(sheet.freeze_panes, 'A2')
        workbook.close()
